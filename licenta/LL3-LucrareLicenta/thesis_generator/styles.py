"""
Modul pentru stilizarea și formatarea documentului Word conform normelor FEAA UCV
Ghid oficial: Anexa 7.1 (A4, TNR 12pt, 1 rând, margini 3cm stânga, 2cm restul, indent 1.25cm)
"""

import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls


def create_base_document():
    doc = docx.Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.different_first_page_header_footer = True
    
    # Configure Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.0
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.space_before = Pt(0)
    
    # Setup Footer with page numbering on following pages
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = f_p.add_run()
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)
    
    return doc


def add_chapter_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text.upper())
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p


def add_subchapter_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p


def add_subsub_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    return p


def add_body_p(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    return p


def add_bullet_p(doc, prefix, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    
    r_bullet = p.add_run("• ")
    r_bullet.font.name = 'Times New Roman'
    r_bullet.font.size = Pt(12)
    r_bullet.font.bold = True
    
    if prefix:
        r_pre = p.add_run(prefix + " ")
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        
    r_text = p.add_run(text)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(12)
    return p


def add_table_with_caption(doc, table_num, title, headers, rows, source="Elaborare proprie a autorului"):
    # Caption deasupra tabelului
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(10)
    p_cap.paragraph_format.space_after = Pt(4)
    p_cap.paragraph_format.first_line_indent = Inches(0)
    r_cap = p_cap.add_run(f"Tabelul {table_num}: {title}")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11)
    r_cap.font.bold = True
    
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Border styling
    tblPr = table._tbl.tblPr
    borders = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''' % nsdecls('w'))
    tblPr.append(borders)
    
    # Headers
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        shading = parse_xml(r'<w:shd %s w:fill="EAEAEA"/>' % nsdecls('w'))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                r.font.bold = True
                
    # Data rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            for p in row_cells[c_idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(10)
                    
    # Sursa subsol
    p_src = doc.add_paragraph()
    p_src.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_src.paragraph_format.space_before = Pt(2)
    p_src.paragraph_format.space_after = Pt(10)
    p_src.paragraph_format.first_line_indent = Inches(0)
    r_src = p_src.add_run(f"Sursa: {source}")
    r_src.font.name = 'Times New Roman'
    r_src.font.size = Pt(10)
    r_src.font.italic = True
    return table


def add_code_snippet(doc, title, code_text, source="Elaborare proprie a autorului în cadrul laboratorului de cercetare"):
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(8)
    p_cap.paragraph_format.space_after = Pt(3)
    p_cap.paragraph_format.first_line_indent = Inches(0)
    r_cap = p_cap.add_run(f"Listarea de cod: {title}")
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(11)
    r_cap.font.bold = True
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shading = parse_xml(r'<w:shd %s w:fill="F5F5F5"/>' % nsdecls('w'))
    cell._tc.get_or_add_tcPr().append(shading)
    
    # Border subtil
    borders = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="none"/>
            <w:left w:val="single" w:sz="12" w:space="0" w:color="003366"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
        </w:tblBorders>
    ''' % nsdecls('w'))
    table._tbl.tblPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run(code_text)
    r.font.name = 'Courier New'
    r.font.size = Pt(9.5)
    
    p_src = doc.add_paragraph()
    p_src.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_src.paragraph_format.space_before = Pt(2)
    p_src.paragraph_format.space_after = Pt(8)
    p_src.paragraph_format.first_line_indent = Inches(0)
    r_src = p_src.add_run(f"Sursa: {source}")
    r_src.font.name = 'Times New Roman'
    r_src.font.size = Pt(9.5)
    r_src.font.italic = True
